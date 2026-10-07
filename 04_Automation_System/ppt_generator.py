from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from reporting_engine import AnalysisResult


BLUE = RGBColor(31, 78, 121)
DARK_BLUE = RGBColor(23, 54, 93)
GREEN = RGBColor(112, 173, 71)
RED = RGBColor(192, 0, 0)
ORANGE = RGBColor(237, 125, 49)
YELLOW = RGBColor(255, 192, 0)
COMPLETED = RGBColor(217, 234, 247)
LIGHT_BG = RGBColor(246, 248, 251)
TEXT = RGBColor(45, 45, 45)
MUTED = RGBColor(100, 100, 100)
WHITE = RGBColor(255, 255, 255)


def ppt_rgb(hex_value: str) -> RGBColor:
    text = str(hex_value).replace("#", "")
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def apply_theme(config: dict[str, Any]) -> None:
    global BLUE, DARK_BLUE, GREEN, RED, ORANGE, YELLOW, COMPLETED, LIGHT_BG, TEXT, MUTED
    theme = config.get("theme", {})
    status = theme.get("status_colors", {})
    BLUE = ppt_rgb(theme.get("primary", "#1F4E78"))
    DARK_BLUE = ppt_rgb(theme.get("secondary", "#305496"))
    GREEN = ppt_rgb(status.get("Active", "#70AD47"))
    RED = ppt_rgb(status.get("Lost", "#C00000"))
    ORANGE = ppt_rgb(status.get("Held Up", "#ED7D31"))
    YELLOW = ppt_rgb(status.get("Yet To Start", "#FFC000"))
    COMPLETED = ppt_rgb(status.get("Completed", "#D9EAF7"))
    LIGHT_BG = ppt_rgb(theme.get("background", "#F7F9FC"))
    TEXT = ppt_rgb(theme.get("text", "#1F1F1F"))
    MUTED = ppt_rgb(theme.get("muted", "#606060"))


def create_powerpoint(result: AnalysisResult, output_path: Path, config: dict[str, Any]) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    apply_theme(config)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs, result, config)
    add_executive_summary_slide(prs, result)
    add_kpi_overview_slide(prs, result)
    add_trend_slide(prs, result)
    add_top_performers_slide(prs, result)
    add_risk_slide(prs, result)
    add_financial_slide(prs, result)
    add_forecast_slide(prs, result)
    add_recommendations_slide(prs, result)
    add_closing_slide(prs, result)

    prs.save(output_path)
    return output_path


def blank_slide(prs: Presentation, title: str, result: AnalysisResult | None = None) -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE

    add_accent_bar(slide)
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.32), Inches(11.6), Inches(0.45))
    frame = title_box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE
    if result:
        add_footer(slide, result)
    return slide


def add_accent_bar(slide: Any) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.11))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.color.rgb = BLUE


def add_footer(slide: Any, result: AnalysisResult) -> None:
    footer = slide.shapes.add_textbox(Inches(0.55), Inches(7.08), Inches(12.2), Inches(0.25))
    p = footer.text_frame.paragraphs[0]
    p.text = f"{result.report_month} {result.report_year} | Source: {result.source.file_path.name}"
    p.font.size = Pt(8.5)
    p.font.color.rgb = MUTED


def add_title_slide(prs: Presentation, result: AnalysisResult, config: dict[str, Any]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = LIGHT_BG

    left_panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(4.2), Inches(7.5))
    left_panel.fill.solid()
    left_panel.fill.fore_color.rgb = BLUE
    left_panel.line.color.rgb = BLUE

    add_metric_pill(slide, Inches(0.65), Inches(0.85), "Monthly Review", WHITE, BLUE)
    title = slide.shapes.add_textbox(Inches(4.65), Inches(1.45), Inches(7.8), Inches(1.0))
    p = title.text_frame.paragraphs[0]
    p.text = "Monthly Executive Review"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    subtitle = slide.shapes.add_textbox(Inches(4.7), Inches(2.45), Inches(7.4), Inches(0.65))
    p = subtitle.text_frame.paragraphs[0]
    p.text = f"{result.report_month} {result.report_year}"
    p.font.size = Pt(22)
    p.font.color.rgb = BLUE

    company = slide.shapes.add_textbox(Inches(4.7), Inches(3.25), Inches(7.4), Inches(0.45))
    p = company.text_frame.paragraphs[0]
    p.text = f"{config.get('company_name', '')} | {config.get('team_name', '')}"
    p.font.size = Pt(14)
    p.font.color.rgb = MUTED

    stats = kpi_map(result)
    add_title_stat(slide, "Total Projects", format_number(stats.get("Total Projects")), 0.75, 2.0)
    add_title_stat(slide, "Completion Rate", format_percent(stats.get("Completion Rate")), 0.75, 3.15)
    add_title_stat(slide, "Ageing Risk", format_number(stats.get("Ageing Risk Count")), 0.75, 4.3)
    add_title_stat(slide, "Avg Cycle Days", format_decimal(stats.get("Average Cycle Days")), 0.75, 5.45)


def add_title_stat(slide: Any, label: str, value: str, x: float, y: float) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(2.95), Inches(0.72))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = value
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = WHITE
    p2 = frame.add_paragraph()
    p2.text = label
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = RGBColor(220, 232, 245)


def add_executive_summary_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "Executive Summary", result)
    stats = kpi_map(result)
    add_kpi_card(slide, 0.65, 1.1, "Projects", format_number(stats.get("Total Projects")), BLUE)
    add_kpi_card(slide, 3.05, 1.1, "Completed", format_number(stats.get("Completed Projects")), GREEN)
    add_kpi_card(slide, 5.45, 1.1, "Loss Rate", format_percent(stats.get("Loss Rate")), RED)
    add_kpi_card(slide, 7.85, 1.1, "SLA Met", format_percent(stats.get("SLA Met Rate")), GREEN)
    add_kpi_card(slide, 10.25, 1.1, "Ageing Risk", format_number(stats.get("Ageing Risk Count")), ORANGE)

    add_bullet_panel(slide, "Key Highlights", result.insights["executive_summary"][:4], 0.75, 2.55, 5.85, 2.15, GREEN)
    add_bullet_panel(slide, "Critical Risks", result.insights["risk_areas"][:4], 6.85, 2.55, 5.75, 2.15, RED)
    add_bullet_panel(slide, "Action Focus", result.insights["recommendations"][:3], 0.75, 5.05, 11.85, 1.45, BLUE)


def add_kpi_overview_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "KPI Overview", result)
    stats = kpi_map(result)
    metrics = [
        ("Revenue", "N/A" if pd.isna(stats.get("Revenue")) else format_currency(stats.get("Revenue")), BLUE),
        ("Profit", "N/A" if pd.isna(stats.get("Profit")) else format_currency(stats.get("Profit")), GREEN),
        ("Growth %", format_percent(stats.get("MoM Volume Growth %")), GREEN if (stats.get("MoM Volume Growth %") or 0) >= 0 else RED),
        ("Conversion", format_percent(stats.get("Conversion Rate")), GREEN),
        ("Utilization", "Owner load", ORANGE),
        ("Productivity", format_number(stats.get("Completed Projects")), GREEN),
        ("SLA Metrics", format_percent(stats.get("SLA Met Rate")), GREEN),
        ("Risk Count", format_number(stats.get("Ageing Risk Count")), RED if (stats.get("Ageing Risk Count") or 0) > 0 else GREEN),
    ]
    for idx, (label, value, color) in enumerate(metrics):
        x = 0.75 + (idx % 4) * 3.05
        y = 1.2 + (idx // 4) * 1.45
        add_kpi_card(slide, x, y, label, value, color)

    add_status_chart(slide, result, 0.85, 4.35, 4.6, 2.25)
    add_bullet_panel(slide, "Operational Reading", result.insights["efficiency_recommendations"][:4], 6.0, 4.15, 6.45, 2.35, BLUE)


def add_trend_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "Trend Analysis", result)
    add_monthly_trend_chart(slide, result, 0.75, 1.05, 7.2, 4.8)
    insights = result.insights["significant_changes"][:2] + result.insights["forecast_commentary"][:2]
    add_bullet_panel(slide, "Observations", insights, 8.25, 1.05, 4.25, 2.35, BLUE)
    add_small_table(slide, result.monthly_trend.tail(6), 8.25, 3.85, 4.25, 2.35, max_rows=6)


def add_top_performers_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "Top Performers", result)
    add_top_performer_chart(slide, result, 0.75, 1.05, 6.2, 4.8)
    table = first_non_empty_table(result.top_performers)
    add_small_table(slide, table, 7.25, 1.1, 5.2, 3.1, max_rows=8)
    add_bullet_panel(slide, "Success Drivers", result.insights["best_performance_areas"][:4], 7.25, 4.55, 5.2, 1.75, GREEN)


def add_risk_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "Risk & Issues", result)
    add_bullet_panel(slide, "Underperforming Areas", result.insights["risk_areas"][:5], 0.75, 1.05, 5.8, 2.35, RED)
    add_bullet_panel(slide, "Root Cause Signals", result.insights["root_cause_analysis"][:5], 6.85, 1.05, 5.75, 2.35, ORANGE)
    add_small_table(slide, result.data_quality, 0.75, 3.85, 11.85, 2.45, max_rows=7)


def add_financial_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "Financial Analysis", result)
    stats = kpi_map(result)
    add_kpi_card(slide, 0.75, 1.05, "Revenue", "N/A" if pd.isna(stats.get("Revenue")) else format_currency(stats.get("Revenue")), BLUE)
    add_kpi_card(slide, 3.45, 1.05, "Cost", "N/A" if pd.isna(stats.get("Cost")) else format_currency(stats.get("Cost")), ORANGE)
    add_kpi_card(slide, 6.15, 1.05, "Profit", "N/A" if pd.isna(stats.get("Profit")) else format_currency(stats.get("Profit")), GREEN)
    add_kpi_card(slide, 8.85, 1.05, "Margin", format_percent(stats.get("Profit Margin")), GREEN)

    if result.financial_columns.get("revenue") is None and result.financial_columns.get("cost") is None:
        notes = [
            "Revenue, cost, profit, and budget variance columns are not present in the current tracker.",
            "Use project volume, cycle time, and employee count as operational proxies this month.",
            "Add commercial fields to the source file to unlock margin and budget variance reporting.",
        ]
    else:
        notes = result.insights["cost_optimization"][:4]
    add_bullet_panel(slide, "Margin & Budget Commentary", notes, 0.75, 2.85, 5.85, 2.5, BLUE)
    proxy_table = result.kpi_summary[result.kpi_summary["KPI"].isin(["Average Cycle Days", "Median Cycle Days", "Total Employee Count", "Average Employee Count"])]
    add_small_table(slide, proxy_table, 7.0, 2.85, 5.45, 2.5, max_rows=5)


def add_forecast_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "Forecast & Outlook", result)
    add_monthly_trend_chart(slide, result, 0.75, 1.05, 6.6, 4.65)
    outlook = result.insights["forecast_commentary"][:3] + result.insights["growth_opportunities"][:3]
    add_bullet_panel(slide, "Forward View", outlook, 7.75, 1.05, 4.85, 2.45, BLUE)
    add_bullet_panel(slide, "Expected Risks", result.insights["risk_areas"][:3], 7.75, 3.95, 4.85, 1.85, ORANGE)


def add_recommendations_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "Recommendations", result)
    add_numbered_actions(slide, result.insights["recommendations"][:6], 1.0, 1.1, 11.2, 5.45)


def add_closing_slide(prs: Presentation, result: AnalysisResult) -> None:
    slide = blank_slide(prs, "Closing Summary", result)
    summary = (
        result.insights["executive_summary"][:2]
        + result.insights["significant_changes"][:1]
        + result.insights["recommendations"][:3]
    )
    add_bullet_panel(slide, "Key Takeaways", summary, 1.0, 1.2, 11.2, 3.3, BLUE)
    add_metric_pill(slide, Inches(1.0), Inches(5.25), "Next monthly run: place latest CSV/XLSX in input folder and run main.py", WHITE, GREEN)


def add_kpi_card(slide: Any, x: float, y: float, label: str, value: str, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.15), Inches(0.95))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(246, 248, 251)
    shape.line.color.rgb = RGBColor(210, 218, 228)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(0.95))
    accent.fill.solid()
    accent.fill.fore_color.rgb = color
    accent.line.color.rgb = color

    text = shape.text_frame
    text.clear()
    p = text.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = value
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE
    p2 = text.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = label
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = MUTED


def add_bullet_panel(
    slide: Any,
    title: str,
    bullets: list[str],
    x: float,
    y: float,
    width: float,
    height: float,
    accent: RGBColor,
) -> None:
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(250, 251, 253)
    panel.line.color.rgb = RGBColor(220, 226, 235)
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(height))
    marker.fill.solid()
    marker.fill.fore_color.rgb = accent
    marker.line.color.rgb = accent

    title_box = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.15), Inches(width - 0.4), Inches(0.3))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    body = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.55), Inches(width - 0.45), Inches(height - 0.65))
    frame = body.text_frame
    frame.word_wrap = True
    frame.clear()
    for idx, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT
        p.space_after = Pt(4)


def add_numbered_actions(slide: Any, actions: list[str], x: float, y: float, width: float, height: float) -> None:
    row_height = height / max(len(actions), 1)
    for idx, action in enumerate(actions, start=1):
        top = y + (idx - 1) * row_height
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(top + 0.05), Inches(0.42), Inches(0.42))
        circle.fill.solid()
        circle.fill.fore_color.rgb = BLUE
        circle.line.color.rgb = BLUE
        p = circle.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = str(idx)
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE

        box = slide.shapes.add_textbox(Inches(x + 0.6), Inches(top), Inches(width - 0.7), Inches(0.58))
        p = box.text_frame.paragraphs[0]
        p.text = action
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT


def add_metric_pill(slide: Any, x: Any, y: Any, text: str, font_color: RGBColor, fill_color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.2), Inches(0.42))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = text
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = font_color


def add_monthly_trend_chart(slide: Any, result: AnalysisResult, x: float, y: float, width: float, height: float) -> None:
    trend = result.monthly_trend.tail(12)
    if trend.empty:
        add_empty_note(slide, "No trend data available", x, y, width, height)
        return
    data = CategoryChartData()
    data.categories = list(trend["Month"])
    data.add_series("Total Projects", list(trend["Total Projects"]))
    data.add_series("Completed", list(trend["Completed"]))
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
        data,
    )
    chart = chart_shape.chart
    for idx, color in enumerate([COMPLETED, GREEN, RED, ORANGE]):
        try:
            point = chart.series[0].points[idx]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = color
        except Exception:
            continue
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.chart_title.text_frame.text = "Monthly Project Trend"
    chart.value_axis.has_major_gridlines = True
    chart.series[0].format.line.color.rgb = BLUE
    chart.series[1].format.line.color.rgb = GREEN


def add_status_chart(slide: Any, result: AnalysisResult, x: float, y: float, width: float, height: float) -> None:
    stats = kpi_map(result)
    data = CategoryChartData()
    data.categories = ["Completed", "Open", "Lost", "Held Up"]
    data.add_series(
        "Projects",
        [
            stats.get("Completed Projects") or 0,
            stats.get("Active/Open Projects") or 0,
            stats.get("Lost Projects") or 0,
            stats.get("Held Up Projects") or 0,
        ],
    )
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(width), Inches(height), data)
    chart = chart_shape.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.BEST_FIT
    chart.chart_title.text_frame.text = "Status Mix"


def add_top_performer_chart(slide: Any, result: AnalysisResult, x: float, y: float, width: float, height: float) -> None:
    if "Assigned Engineer" in result.top_performers and not result.top_performers["Assigned Engineer"].empty:
        table = result.top_performers["Assigned Engineer"].head(8)
        category_col = "Assigned Engineer"
    else:
        table = first_non_empty_table(result.top_performers).head(8)
        category_col = table.columns[0] if not table.empty else None
    if table.empty or not category_col:
        add_empty_note(slide, "No performer data available", x, y, width, height)
        return
    data = CategoryChartData()
    data.categories = list(table[category_col].astype(str))
    data.add_series("Completed", list(table["Completed"]))
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(width), Inches(height), data)
    chart = chart_shape.chart
    chart.has_legend = False
    chart.chart_title.text_frame.text = f"Top {category_col} by Completions"
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = GREEN


def add_small_table(slide: Any, df: pd.DataFrame, x: float, y: float, width: float, height: float, max_rows: int = 6) -> None:
    if df is None or df.empty:
        add_empty_note(slide, "No table data available", x, y, width, height)
        return
    data = df.head(max_rows).copy()
    columns = list(data.columns[:5])
    data = data[columns]
    rows = len(data) + 1
    cols = len(columns)
    table_shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(width), Inches(height))
    table = table_shape.table
    for idx, column in enumerate(columns):
        cell = table.cell(0, idx)
        cell.text = str(column)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(8.5)
            paragraph.font.bold = True
            paragraph.font.color.rgb = WHITE
    for row_idx, row in enumerate(data.itertuples(index=False), start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = format_table_value(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(7.8)
                paragraph.font.color.rgb = TEXT


def add_empty_note(slide: Any, text: str, x: float, y: float, width: float, height: float) -> None:
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(250, 251, 253)
    box.line.color.rgb = RGBColor(220, 226, 235)
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(13)
    p.font.color.rgb = MUTED


def first_non_empty_table(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    for table in tables.values():
        if not table.empty:
            return table
    return pd.DataFrame()


def kpi_map(result: AnalysisResult) -> dict[str, Any]:
    return dict(zip(result.kpi_summary["KPI"], result.kpi_summary["Value"]))


def format_number(value: Any) -> str:
    if value is None or pd.isna(value):
        return "0"
    return f"{float(value):,.0f}"


def format_decimal(value: Any) -> str:
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):,.1f}"


def format_percent(value: Any) -> str:
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):.1%}"


def format_currency(value: Any) -> str:
    if value is None or pd.isna(value):
        return "N/A"
    return f"INR {float(value):,.0f}"


def format_table_value(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, float):
        if abs(value) <= 1:
            return f"{value:.1%}"
        return f"{value:,.1f}"
    return str(value)
